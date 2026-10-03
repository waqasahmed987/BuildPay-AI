"""Deterministic tools used by the agents.
All quantity and money arithmetic lives here - never in the LLM."""
import io
import json
import os

import streamlit as st

from config import BOQ_FILES, PLATFORM_RULES, REQUIRED_DOCS
from db import ex, now, q


# ---------- BOQ tool ----------
def load_boq_file(size):
    """BOQ Excel Loader: read item code, section, description, unit, quantity and rate from the template."""
    from openpyxl import load_workbook
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), BOQ_FILES[size])
    if not os.path.exists(path):
        raise FileNotFoundError(f"BOQ file '{BOQ_FILES[size]}' not found - upload it to GitHub next to app.py.")
    ws = load_workbook(path, data_only=True)["BOQ"]
    header, items = None, []
    for row in ws.iter_rows(values_only=True):
        if header is None:  # find the table header row
            if row and row[0] == "Item Code":
                header = {str(n).strip(): i for i, n in enumerate(row) if n}
            continue
        if not row[0]:      # blank / TOTAL row
            continue
        get = lambda col: row[header[col]]
        items.append({"code": str(get("Item Code")).strip(), "section": get("Section"),
                      "description": get("Description"), "unit": get("Unit"),
                      "qty": float(get("Qty")), "rate": float(get("Unit Rate (PKR)"))})
    if not items:
        raise ValueError(f"No BOQ items found in {BOQ_FILES[size]}.")
    return items


# ---------- Document tools ----------
def read_document(f):
    """Document Reader: extract text from an uploaded file."""
    name, data = f.name.lower(), f.getvalue()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        elif name.endswith(".docx"):
            from docx import Document
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            text = "\n".join("\t".join(str(c) for c in row if c is not None)
                             for ws in wb for row in ws.iter_rows(values_only=True))
        else:
            text = data.decode("utf-8", errors="ignore")
    except Exception:
        text = ""
    return text[:20000]


def save_documents(entity, entity_id, files):
    for f in files or []:
        ex("INSERT INTO docs(entity,entity_id,filename,text,uploaded) VALUES(?,?,?,?,?)",
           (entity, entity_id, f.name, read_document(f), now()))


def document_excerpts(cr_id, limit=3, chars=600):
    docs = q("SELECT filename, text FROM docs WHERE entity='CR' AND entity_id=?", (cr_id,))
    return {d["filename"]: (d["text"][:chars] or "(no readable text)") for d in docs[:limit]}


def required_document_checker(cr_id):
    """Required Document Checker. Keyword match only - NOT verification."""
    cr = q("SELECT b.section FROM crs c JOIN boq b ON b.id=c.boq_id WHERE c.id=?", (cr_id,))[0]
    docs = q("SELECT filename, text FROM docs WHERE entity='CR' AND entity_id=?", (cr_id,))
    blobs = [(d["filename"] + " " + d["text"][:5000]).lower() for d in docs]
    found, missing = [], []
    for label, kws in REQUIRED_DOCS[cr["section"]]:
        (found if any(k in b for b in blobs for k in kws) else missing).append(label)
    return {"uploaded_files": [d["filename"] for d in docs],
            "keyword_matched_required_docs": found,
            "missing_required_docs": missing,
            "files_with_no_readable_text": [d["filename"] for d in docs if not d["text"].strip()],
            "note": "Upload or keyword match does not mean the evidence is verified."}


# ---------- Lookup tools ----------
def cr_lookup(cr_id):
    return q("SELECT c.*, b.code, b.section, b.description AS item, b.unit, b.rate "
             "FROM crs c JOIN boq b ON b.id=c.boq_id WHERE c.id=?", (cr_id,))[0]


def boq_lookup(boq_id):
    """BOQ Lookup: exact activity, quantity, unit and rate."""
    return q("SELECT * FROM boq WHERE id=?", (boq_id,))[0]


def get_contract(project_id):
    return (q("SELECT contract FROM projects WHERE id=?", (project_id,))[0]["contract"] or "").strip()


def latest_measurement_cr(boq_id):
    rows = q("SELECT cr_id FROM measurements WHERE boq_id=? ORDER BY id DESC LIMIT 1", (boq_id,))
    return rows[0]["cr_id"] if rows else None


# ---------- Quantity tools ----------
def _ipc_qty(project_id, boq_id, statuses):
    ph = ",".join("?" * len(statuses))
    rows = q(f"SELECT lines_json FROM ipcs WHERE project_id=? AND status IN ({ph})", (project_id, *statuses))
    return sum(l["current_qty"] for r in rows for l in json.loads(r["lines_json"]) if l["boq_id"] == boq_id)


def quantity_calculator(boq_id):
    """Quantity Calculator: original, approved, measured, certified and overrun quantities."""
    it = boq_lookup(boq_id)
    var = q("SELECT COALESCE(SUM(additional),0) s FROM variations WHERE boq_id=? AND status='Approved'", (boq_id,))[0]["s"]
    measured = q("SELECT COALESCE(SUM(qty),0) s FROM measurements WHERE boq_id=?", (boq_id,))[0]["s"]
    certified = _ipc_qty(it["project_id"], boq_id, ("Approved",))
    committed = _ipc_qty(it["project_id"], boq_id, ("Draft", "Approved"))
    approved = it["qty"] + var
    r = lambda x: round(x, 3)
    return {"boq_id": boq_id, "code": it["code"], "description": it["description"], "unit": it["unit"],
            "rate": it["rate"], "original_qty": r(it["qty"]), "approved_variation_qty": r(var),
            "approved_qty": r(approved), "measured_cumulative": r(measured),
            "certified_previous": r(certified), "in_draft_ipc": r(committed - certified),
            "unbilled": r(measured - committed), "remaining_approved": r(approved - measured),
            "overrun": r(max(0, measured - approved))}


def pending_variation_qty(boq_id):
    return q("SELECT COALESCE(SUM(additional),0) s FROM variations WHERE boq_id=? AND status='Submitted'", (boq_id,))[0]["s"]


def approved_cr_quantity(boq_id, exclude_cr_id=None):
    return q("SELECT COALESCE(SUM(req_qty),0) s FROM crs WHERE boq_id=? AND status='Approved' AND id!=?",
             (boq_id, exclude_cr_id or -1))[0]["s"]


def historical_record_query(boq_id, cr_id, req_qty):
    """Historical Record Query + duplicate detection."""
    prior = q("SELECT cr_no, req_qty, status FROM crs WHERE boq_id=? AND id!=?", (boq_id, cr_id))
    dups = [p["cr_no"] for p in prior if p["status"] != "Rejected" and abs(p["req_qty"] - req_qty) < 1e-9]
    variations = q("SELECT var_no, additional, status FROM variations WHERE boq_id=?", (boq_id,))
    return {"prior_crs_same_item": prior, "possible_duplicate_crs_same_quantity": dups,
            "variations_same_item": variations,
            "certified_previous": quantity_calculator(boq_id)["certified_previous"]}


# ---------- IPC tools ----------
def build_ipc_lines(project_id):
    """Only approved quantity is eligible; excess is flagged, never silently certified."""
    lines, exceptions = [], []
    for it in q("SELECT id FROM boq WHERE project_id=? ORDER BY id", (project_id,)):
        s = quantity_calculator(it["id"])
        if s["unbilled"] <= 0:
            continue
        prev = s["certified_previous"] + s["in_draft_ipc"]
        eligible = round(max(0, min(s["unbilled"], s["approved_qty"] - prev)), 3)
        excess = round(s["unbilled"] - eligible, 3)
        if excess > 0:
            exceptions.append(f"{s['code']}: {excess} {s['unit']} measured beyond approved quantity - "
                              "approved variation required; excluded from this IPC.")
        if eligible <= 0:
            continue
        crs = q("SELECT DISTINCT c.id, c.cr_no FROM measurements m JOIN crs c ON c.id=m.cr_id WHERE m.boq_id=?", (s["boq_id"],))
        for c in crs:
            missing = required_document_checker(c["id"])["missing_required_docs"]
            if missing:
                exceptions.append(f"{c['cr_no']} ({s['code']}): missing evidence - {', '.join(missing)}.")
        vars_ = [v["var_no"] for v in q("SELECT var_no FROM variations WHERE boq_id=? AND status='Approved'", (s["boq_id"],))]
        lines.append({"boq_id": s["boq_id"], "code": s["code"], "description": s["description"], "unit": s["unit"],
                      "rate": s["rate"], "previous_qty": round(prev, 3), "current_qty": eligible,
                      "cumulative_qty": round(prev + eligible, 3), "amount": round(eligible * s["rate"], 2),
                      "source_crs": ", ".join(c["cr_no"] for c in crs), "approved_variations": ", ".join(vars_)})
    return lines, exceptions


def payment_calculator(lines, retention_pct, other_deductions):
    """Payment Calculator: gross, retention and net as deterministic arithmetic."""
    gross = round(sum(l["amount"] for l in lines), 2)
    retention = round(gross * retention_pct / 100, 2)
    return {"gross": gross, "retention_pct": retention_pct, "retention": retention,
            "other_deductions": other_deductions, "net": round(gross - retention - other_deductions, 2)}


# ---------- Audit + content tools ----------
def audit_recorder(event, entity, entity_id, detail=""):
    """Audit Recorder: append-only record of a material workflow event."""
    ex("INSERT INTO audit(ts,role,actor,event,entity,entity_id,detail) VALUES(?,?,?,?,?,?,?)",
       (now(), st.session_state.get("role", "-"), st.session_state.get("person") or "-",
        event, entity, entity_id, detail))
    return "recorded"


def platform_rules(platform):
    return PLATFORM_RULES.get(platform, "Keep it clear and concise.")
